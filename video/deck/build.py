"""Builds out/simpuru-deck.pptx from the Slides composition, with the film embedded on slide 2.

    npx remotion still Slides out/slides/slide-NN.png --frame=N   (for every slide, see README)
    python3 deck/build.py
"""
import glob

from pptx import Presentation
from pptx.util import Emu

W, H = 12192000, 6858000  # 16:9
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
slides = sorted(glob.glob("out/slides/slide-*.png"))
for i, png in enumerate(slides):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(png, 0, 0, W, H)
    if i == 1:  # the demo plays inside the deck (no external links)
        pad = Emu(int(W * 0.06))
        vw = W - 2 * pad
        vh = int(vw * 9 / 16)
        s.shapes.add_movie("out/simpuru-demo.mp4", pad, int((H - vh) / 2) - Emu(int(H * 0.02)), vw, vh,
                           poster_frame_image="public/deck/f528.png", mime_type="video/mp4")
prs.save("out/simpuru-deck.pptx")
print("saved out/simpuru-deck.pptx,", len(slides), "slides")
